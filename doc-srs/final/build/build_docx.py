"""Build the final Laporan Capstone docx from the markdown draft.

Base: doc-srs/final/build/template-dasar/Laporan Capstone Project Kelompok 1.docx (official template, cover already filled).
Content: doc-srs/draft/Laporan-Capstone-SAPA-draft.md + mockup spec + VP diagram PNGs.
Citations: Word-native CITATION / BIBLIOGRAPHY fields backed by customXml sources (referensi-sitasi/refs.json).

Usage: python build_docx.py   (run with a python that has python-docx)
"""
import copy
import json
import os
import re
from xml.sax.saxutils import escape

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))  # doc-srs
BASE = os.path.join(ROOT, "final", "build", "template-dasar", "Laporan Capstone Project Kelompok 1.docx")
DRAFT = os.path.join(ROOT, "draft", "Laporan-Capstone-SAPA-draft.md")
MOCKSPEC = os.path.join(ROOT, "draft", "mockup-spec-untuk-figma.md")
PNG = os.path.join(ROOT, "diagram", "png")
MOCK = os.path.join(ROOT, "final", "mockup")
ASSETS = os.path.join(ROOT, "assets")
LOGO = os.path.join(ASSETS, "logo", "logo-umb.png")
XP_FIG = os.path.join(ASSETS, "gambar", "sdlc-extreme-programming.png")  # Gambar siklus XP (SDLC)
REFS = json.load(open(os.path.join(ROOT, "referensi-sitasi", "refs.json"), encoding="utf-8"))
OUT = os.path.join(ROOT, "final", "Laporan Capstone Project Kelompok 1 - Final.docx")

MAX_W, MAX_H = Cm(15.0), Cm(21.5)

doc = docx.Document(BASE)
body = doc.element.body
kids = list(body)

# ------------------------------------------------------------------ keep cover + front matter skeleton
xp_blip = None
for k in kids:
    if "Siklus Hidup" in "".join(t.text or "" for t in k.iter(qn("w:t"))):
        break
    blips = list(k.iter(qn("a:blip")))
    if blips:
        xp_blip = blips[0].get(qn("r:embed"))
xp_image = doc.part.related_parts[xp_blip].blob

def text_of(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))

keep_upto = None
for i, k in enumerate(kids):
    if k.tag == qn("w:sdt"):  # TOC content control right after DAFTAR ISI
        keep_upto = i
        break
final_sect = kids[-1]
front_sect = None
for k in kids[keep_upto + 1:]:
    if k.tag == qn("w:p") and k.find(".//" + qn("w:sectPr")) is not None and "BAB I" not in text_of(k):
        front_sect = k  # section break between front matter and BAB I (page numbering restart)
for k in kids[keep_upto + 1:]:
    if k is not final_sect and k is not front_sect:
        body.remove(k)

# cover: full names and subtitle
NAMES = {"Rifqi Adhi S.": "Rifqi Adhi Saputra", "Isnaini Anggista D.": "Isnaini Anggista Damayanti",
         "Enrico Yuda D.": "Enrico Yuda Dwinata", "Yordan Gabriel": "Yordan Gabriell",
         "Agung Nugraha E.S": "Agung Nugraha Eka Saputra"}
for tblel in body.iter(qn("w:tbl")):
    for pel in tblel.iter(qn("w:p")):
        ts = list(pel.iter(qn("w:t")))
        joined = "".join(t.text or "" for t in ts).strip()
        if joined in NAMES and ts:
            ts[0].text = NAMES[joined]
            for t in ts[1:]:
                t.text = ""
# cover logo: replace the template logo with assets/logo/logo-umb.png (falls back to the template image if missing)
def swap_cover_logo():
    if not os.path.exists(LOGO):
        print("logo: file tidak ditemukan, memakai logo template:", LOGO)
        return
    from docx.image.image import Image as _Img
    img = _Img.from_file(LOGO)
    for k in list(body)[:keep_upto]:
        for blip in k.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            if not rid or not doc.part.related_parts[rid].partname.endswith("image1.jpeg"):
                continue
            # keep the template's image part (same name/content type) and swap its bytes for the new logo (as JPEG on white)
            import io
            from PIL import Image as _PIL
            pil = _PIL.open(LOGO).convert("RGBA")
            bg = _PIL.new("RGB", pil.size, (255, 255, 255)); bg.paste(pil, mask=pil.split()[3])
            buf = io.BytesIO(); bg.save(buf, "JPEG", quality=95, dpi=(300, 300))
            doc.part.related_parts[rid]._blob = buf.getvalue()
            parent = blip.getparent()
            for sr in parent.findall(qn("a:srcRect")):
                parent.remove(sr)
            drawing = blip
            while drawing is not None and drawing.tag not in (qn("wp:inline"), qn("wp:anchor")):
                drawing = drawing.getparent()
            ext = drawing.find(qn("wp:extent"))
            cy = int(ext.get("cy"))
            cx = int(round(cy * img.px_width / img.px_height))
            ext.set("cx", str(cx)); ext.set("cy", str(cy))
            for a_ext in drawing.iter(qn("a:ext")):
                if a_ext.get("cx") is not None:
                    a_ext.set("cx", str(cx)); a_ext.set("cy", str(cy))
            print("logo:", os.path.relpath(LOGO, ROOT))
            return
swap_cover_logo()
# cover lines use heading styles; keep them out of the table of contents
for k in list(body)[:keep_upto]:
    if k.tag == qn("w:p"):
        pPr = k.find(qn("w:pPr"))
        st = pPr.find(qn("w:pStyle")) if pPr is not None else None
        if st is not None and st.get(qn("w:val")) in ("Judul1", "Judul2"):
            size = "32" if st.get(qn("w:val")) == "Judul1" else "28"
            st.set(qn("w:val"), "Normal")
            jc = pPr.find(qn("w:jc"))
            if jc is None:
                jc = OxmlElement("w:jc"); pPr.append(jc)
            jc.set(qn("w:val"), "center")
            for r in k.iter(qn("w:r")):
                rPr = r.find(qn("w:rPr"))
                if rPr is None:
                    rPr = OxmlElement("w:rPr"); r.insert(0, rPr)
                rPr.append(OxmlElement("w:b"))
                sz = OxmlElement("w:sz"); sz.set(qn("w:val"), size); rPr.append(sz)
for p in doc.paragraphs[:6]:
    if p.text.startswith("SAPA"):
        r = p.add_run()
        r.add_break()
        r2 = p.add_run("Chatbot WhatsApp Omnichannel Berbasis Kecerdasan Artifisial dan CRM pada Bimbingan Belajar Asentra")
        r2.font.size = Pt(12)
        r2.bold = False

# ------------------------------------------------------------------ helpers
cursor = front_sect  # new content is inserted before this element until we pass the front matter


def insert(el):
    cursor.addprevious(el)
    return el


def new_par(style=None, align=None):
    p = OxmlElement("w:p")
    insert(p)
    para = docx.text.paragraph.Paragraph(p, doc._body)
    if style:
        para.style = doc.styles[style]
    if align is not None:
        para.alignment = align
    return para


def page_break_before(para):
    pPr = para._p.get_or_add_pPr()
    el = OxmlElement("w:pageBreakBefore")
    pPr.insert(0, el)


def keep_next(para):
    pPr = para._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:keepNext"))


def field(para, instr, result="", bold=False, size=None):
    def run():
        r = OxmlElement("w:r")
        if bold or size:
            rPr = OxmlElement("w:rPr")
            if bold:
                rPr.append(OxmlElement("w:b"))
            if size:
                sz = OxmlElement("w:sz")
                sz.set(qn("w:val"), str(size * 2))
                rPr.append(sz)
            r.append(rPr)
        para._p.append(r)
        return r
    r = run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r.append(fc)
    r = run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r.append(it)
    r = run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r.append(fc)
    r = run(); t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = result; r.append(t)
    r = run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r.append(fc)


# citation lookup: (surname, year) -> tag
CITE = {}
for r in REFS:
    first = r["authors"][0][0] if r.get("authors") else "Undang-Undang"
    CITE[(first.split()[0].replace("González-Flores", "González-Flores"), r["year"])] = r["tag"]
CITE[("Undang-Undang", 2022)] = "UUPDP2022"
LANG = 1057  # Indonesian

PAREN = re.compile(r"\(([^()]*?(?:19|20)\d{2})\)")
NARR = re.compile(r"\b([A-Z][\wà-ÿ\-]+)((?: dkk\.)|(?: dan [A-Z][\wà-ÿ\-]+))? \(((?:19|20)\d{2})\)")


def cite_tags(inner):
    tags = []
    for part in inner.split(";"):
        part = part.strip()
        m = re.search(r"((?:19|20)\d{2})$", part)
        if not m:
            return None
        year = int(m.group(1))
        surname = part.split()[0].rstrip(",")
        if part.startswith("Undang-Undang"):
            surname = "Undang-Undang"
        tag = CITE.get((surname, year))
        if not tag:
            return None
        tags.append(tag)
    return tags


def add_text(para, text, bold=False, italic=False, size=None):
    """Write text with **bold**/*italic* markup, citations as Word fields, [ISI] highlighted."""
    text = text.replace("`", "")
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[ISI[^\]]*\])", text)
    for tok in tokens:
        if not tok:
            continue
        b, i = bold, italic
        if tok.startswith("**") and tok.endswith("**"):
            tok, b = tok[2:-2], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, i = tok[1:-1], True
        if tok.startswith("[ISI"):
            r = para.add_run(tok)
            r.bold, r.italic = b, i
            if size: r.font.size = Pt(size)
            rPr = r._r.get_or_add_rPr(); hl = OxmlElement("w:highlight"); hl.set(qn("w:val"), "yellow"); rPr.append(hl)
            continue
        add_cited(para, tok, b, i, size)


def add_cited(para, text, b, i, size):
    pos = 0
    events = []
    for m in PAREN.finditer(text):
        tags = cite_tags(m.group(1))
        if tags:
            events.append((m.start(), m.end(), "paren", tags, m.group(0)))
    for m in NARR.finditer(text):
        if any(s <= m.start() < e for s, e, *_ in events):
            continue
        key = (m.group(1), int(m.group(3)))
        if key in CITE:
            yr_start = m.end() - 6
            events.append((yr_start, m.end(), "narr", [CITE[key]], m.group(0)[-6:]))
    events.sort()
    for s, e, kind, tags, shown in events:
        if s > pos:
            r = para.add_run(text[pos:s]); r.bold, r.italic = b, i
            if size: r.font.size = Pt(size)
        instr = f"CITATION {tags[0]} \\l {LANG} " + ("\\n " if kind == "narr" else "") + "".join(f"\\m {t} " for t in tags[1:])
        field(para, instr.strip(), shown, size=size)
        pos = e
    if pos < len(text):
        r = para.add_run(text[pos:]); r.bold, r.italic = b, i
        if size: r.font.size = Pt(size)


def heading(level, text, pagebreak=False):
    style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    p = new_par(style)
    p.add_run(text)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if pagebreak:
        page_break_before(p)
    keep_next(p)
    return p


def body_par(text, italic=False, indent=True):
    p = new_par("Body Text", WD_ALIGN_PARAGRAPH.JUSTIFY)
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    if indent:
        pf.first_line_indent = Cm(1.0)
    add_text(p, text, italic=italic)
    return p


def list_par(marker, text):
    p = new_par("List Paragraph", WD_ALIGN_PARAGRAPH.JUSTIFY)
    pf = p.paragraph_format
    pf.left_indent, pf.first_line_indent = Cm(1.0), Cm(-0.6)
    pf.line_spacing = 1.5
    pf.space_after = Pt(2)
    p.add_run(marker + "\t")
    tabs = pf.tab_stops; tabs.add_tab_stop(Cm(1.0))
    add_text(p, text)
    return p


SEQ = {"Tabel": 0, "Gambar": 0}


def caption(kind, text):
    SEQ[kind] += 1
    p = new_par("Caption", WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(f"{kind} 3.")
    field(p, f"SEQ {kind} \\* ARABIC", str(SEQ[kind]))
    p.add_run(" " + text)
    if kind == "Tabel":
        keep_next(p)
    return p


def set_cell_borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), "5000"); w.set(qn("w:type"), "pct")
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    tblPr.append(w)


def shade(cell, color="D9D9D9"):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), color)
    tcPr.append(s)


def cell_text(cell, text, bold=False, align=None):
    p = cell.paragraphs[0]
    p.style = doc.styles["Normal"]
    p.paragraph_format.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    add_text(p, text, bold=bold, size=10)


def new_table(rows, cols, widths=None):
    tbl = doc.add_table(rows=rows, cols=cols)
    insert(tbl._tbl)
    set_cell_borders(tbl)
    if widths:
        for row in tbl.rows:
            for c, w in zip(row.cells, widths):
                c.width = Cm(w)
        grid = tbl._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w * 567)))
    return tbl


def md_table(header, rows):
    n = len(header)
    # column widths proportional to content length, short columns (No, Kode) kept narrow
    need = []
    for j in range(n):
        longest = max([len(header[j])] + [len(r[j]) if j < len(r) else 0 for r in rows])
        avg = sum(len(r[j]) if j < len(r) else 0 for r in rows) / max(1, len(rows))
        need.append(max(1.1, min(longest, 12) * 0.22 if longest <= 6 else 2.2 + avg * 0.06))
    tot = sum(need)
    widths = [round(15.0 * x / tot, 2) for x in need]
    tbl = new_table(1 + len(rows), n, widths)
    for j, h in enumerate(header):
        cell_text(tbl.rows[0].cells[j], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(tbl.rows[0].cells[j])
    hdr = tbl.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); hdr.append(th)
    for i, r in enumerate(rows, 1):
        for j in range(n):
            cell_text(tbl.rows[i].cells[j], r[j] if j < len(r) else "")
    spacer()
    return tbl


def spacer():
    p = new_par("Normal")
    p.paragraph_format.space_after = Pt(4)


def uc_table(elem, flow):
    """ANSI use case description: label column + Aktor | Sistem basic flow."""
    order = ["Nama Use Case", "Nomor Use Case", "Aktor", "Deskripsi Singkat", "Pre Condition"]
    after = ["Post Condition", "Alternative Flow", "Error Condition"]
    rows = len(order) + 1 + len(flow) + len(after)
    tbl = new_table(rows, 3, [3.6, 5.7, 5.7])
    r = 0

    def label_row(label, value):
        nonlocal r
        row = tbl.rows[r]
        cell_text(row.cells[0], label, bold=True); shade(row.cells[0], "EDEDED")
        merged = row.cells[1].merge(row.cells[2])
        cell_text(merged, value)
        r += 1
    for k in order:
        label_row(k, elem.get(k, "-"))
    first = r
    hdr = tbl.rows[r]
    cell_text(hdr.cells[1], "Aktor", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER); shade(hdr.cells[1])
    cell_text(hdr.cells[2], "Sistem", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER); shade(hdr.cells[2])
    r += 1
    for a, s in flow:
        cell_text(tbl.rows[r].cells[1], a)
        cell_text(tbl.rows[r].cells[2], s)
        r += 1
    lab = tbl.cell(first, 0).merge(tbl.cell(r - 1, 0))
    cell_text(lab, "Basic Flow of Event", bold=True); shade(lab, "EDEDED")
    for k in after:
        label_row(k, elem.get(k, "-"))
    spacer()


def picture(path_or_blob, desc):
    p = new_par("Normal", WD_ALIGN_PARAGRAPH.CENTER)
    keep_next(p)
    run = p.add_run()
    import io
    from docx.image.image import Image
    if isinstance(path_or_blob, bytes):
        img = Image.from_blob(path_or_blob); src = io.BytesIO(path_or_blob)
    else:
        img = Image.from_file(path_or_blob); src = path_or_blob
    w, h = img.px_width, img.px_height
    scale = min(MAX_W / w, MAX_H / h)
    run.add_picture(src, width=int(w * scale))
    caption("Gambar", desc)


def placeholder(desc, fname):
    tbl = new_table(1, 1, [15.0])
    c = tbl.rows[0].cells[0]
    cell_text(c, f"[ISI: mockup high-fidelity dari Figma, simpan sebagai final/mockup/{fname}, lalu jalankan ulang build]",
              align=WD_ALIGN_PARAGRAPH.CENTER)
    tr = tbl.rows[0]._tr.get_or_add_trPr(); hh = OxmlElement("w:trHeight"); hh.set(qn("w:val"), "3400"); tr.append(hh)
    caption("Gambar", desc)


# ------------------------------------------------------------------ front matter lists
def front_list(title, instr):
    p = new_par("Heading 1", WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run(title)
    page_break_before(p)
    q = new_par("table of figures")
    field(q, instr, "(Klik kanan → Update Field untuk memperbarui daftar ini)")


front_list("DAFTAR TABEL", 'TOC \\h \\z \\c "Tabel"')
front_list("DAFTAR GAMBAR", 'TOC \\h \\z \\c "Gambar"')

# everything after the front-matter section break
anchor = OxmlElement("w:p")
final_sect.addprevious(anchor)
cursor = anchor

# ------------------------------------------------------------------ parse draft
md = open(DRAFT, encoding="utf-8").read()
start = md.index("## BAB I PENDAHULUAN")
lines = md[start:].split("\n")

mock_desc = {}
for m in re.finditer(r"- \*\*(MK-\d+)\*\*: (.+)", open(MOCKSPEC, encoding="utf-8").read()):
    mock_desc[m.group(1)] = m.group(2).strip()
MOCKS = [("MK-01", "MK-01-login.png", "Login"), ("MK-02", "MK-02-dasbor.png", "Dasbor Pemantauan"),
         ("MK-03", "MK-03-kotak-masuk.png", "Kotak Masuk"), ("MK-04", "MK-04-prospek.png", "Prospek"),
         ("MK-05", "MK-05-kontak.png", "Kontak"), ("MK-06", "MK-06-nomor-whatsapp.png", "Nomor WhatsApp"),
         ("MK-07", "MK-07-basis-pengetahuan.png", "Basis Pengetahuan"),
         ("MK-08", "MK-08-pengaturan-ai.png", "Pengaturan AI dan Konsol Uji"),
         ("MK-09", "MK-09-antrean-risiko.png", "Antrean Keluar dan Pusat Risiko"),
         ("MK-10", "MK-10-tim-akses.png", "Tim dan Akses"), ("MK-11", "MK-11-log-audit.png", "Log Audit")]

UC_NAMES = ["Login", "Mengirim Pesan WhatsApp", "Menjawab Pesan Otomatis dengan AI", "Mengelola Percakapan",
            "Mengelola Prospek", "Mengelola Kontak Pelanggan", "Melihat Dasbor Pemantauan",
            "Mengelola Nomor WhatsApp", "Mengelola Basis Pengetahuan", "Mengatur AI dan Anti-ban",
            "Mengelola Tim dan Hak Akses", "Melihat Log Audit"]


def png(name):
    return os.path.join(PNG, name + ".png")


def gen_341():
    body_par("Use case diagram SAPA menggambarkan interaksi empat aktor utama (Pelanggan, Agen/CS, Supervisor, "
             "dan Administrator) serta dua aktor sekunder (WhatsApp Web dan Penyedia AI) dengan dua belas use case. "
             "Setiap use case yang memerlukan autentikasi memiliki relasi <<Include>> ke use case Login, sedangkan "
             "Mengirim Pesan WhatsApp tidak memerlukan login karena dilakukan pelanggan melalui aplikasi WhatsApp dan "
             "meng-include Menjawab Pesan Otomatis dengan AI yang dijalankan sistem.")
    picture(png("Use Case Diagram SAPA"), "Use Case Diagram SAPA")


def gen_343():
    body_par("Sequence diagram disusun untuk setiap use case pada use case diagram. Setiap diagram memuat aktor, "
             "objek antarmuka (View/boundary), pengendali (Cont/control), dan entitas (Ent/entity) beserta urutan "
             "pesan bernomor sesuai Basic Flow of Event pada use case description.")
    for i, n in enumerate(UC_NAMES, 1):
        picture(png(f"SD-{i:02d} - {n}"), f"Sequence Diagram SD-{i:02d} {n}")


def gen_344():
    body_par("Class diagram disusun dengan tata letak Model-View-Controller (MVC). Kolom kiri berisi kelas antarmuka "
             "(UI) beserta atribut dan operasi get/set, kolom tengah berisi kelas pengendali (Cont) yang memuat logika "
             "proses setiap use case, dan kolom kanan berisi kelas entitas (Ent) yang merepresentasikan data tersimpan.")
    picture(png("Class Diagram SAPA (MVC)"), "Class Diagram SAPA dengan Layout MVC")


def gen_345():
    body_par("Activity diagram diturunkan dari setiap use case description. Swimlane kiri memuat aktivitas aktor dan "
             "swimlane kanan memuat aktivitas sistem; percabangan (decision) menggambarkan Alternative Flow dan "
             "Error Condition.")
    for i, n in enumerate(UC_NAMES, 1):
        picture(png(f"AD-{i:02d} - {n}"), f"Activity Diagram AD-{i:02d} {n}")


def gen_35():
    body_par("Mockup aplikasi dibuat dalam bentuk high-fidelity menggunakan Figma untuk setiap halaman yang akan "
             "dibangun. Setiap halaman dipetakan ke use case yang dilayaninya.")
    for code, fname, title in MOCKS:
        p = os.path.join(MOCK, fname)
        sub = new_par("Body Text", WD_ALIGN_PARAGRAPH.JUSTIFY)
        sub.paragraph_format.line_spacing = 1.5
        add_text(sub, f"**{code}: Halaman {title}.** {mock_desc.get(code, '')}")
        keep_next(sub)
        if os.path.exists(p):
            picture(p, f"Mockup Halaman {title}")
        else:
            placeholder(f"Mockup Halaman {title}", fname)


GENERATED = {"3.4.1": gen_341, "3.4.3": gen_343, "3.4.4": gen_344, "3.4.5": gen_345, "3.5": gen_35}

pending_caption = None
pending_elem = None  # UC element table waiting for its flow table
skip = False
in_biblio = False
biblio = []
i = 0


def parse_table(start):
    rows = []
    j = start
    while j < len(lines) and lines[j].startswith("|"):
        cells = [c.strip() for c in lines[j].strip().strip("|").split("|")]
        rows.append(cells)
        j += 1
    rows = [r for r in rows if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
    return rows, j


while i < len(lines):
    ln = lines[i]
    s = ln.strip()
    if s.startswith("#"):
        level = len(s) - len(s.lstrip("#"))
        title = s.lstrip("#").strip()
        skip = False
        if title.startswith("DAFTAR PUSTAKA"):
            heading(1, "DAFTAR PUSTAKA", pagebreak=True)
            in_biblio = True
            i += 1
            continue
        if level == 2:
            heading(1, title, pagebreak=True)
        elif level == 3:
            heading(2, title)
        else:
            heading(3, title)
        num = title.split()[0]
        if num in GENERATED:
            GENERATED[num]()
            skip = num != "3.4.2"
        if num == "3.4.2":
            body_par("Use case description disusun berdasarkan use case diagram pada Gambar 3.1. Setiap tabel memuat "
                     "nama, nomor, aktor, deskripsi singkat, kondisi awal, alur utama (aktor dan sistem), kondisi akhir, "
                     "alur alternatif, dan kondisi galat.")
            skip = False
        i += 1
        continue
    if skip:
        i += 1
        continue
    if in_biblio:
        if s and not s.startswith("*("):
            biblio.append(re.sub(r"\s*\*\(bawaan template\)\*", "", s))
        i += 1
        continue
    if not s or s == "---" or s.startswith(">") or s.startswith("*(") or s.startswith("!["):
        i += 1
        continue
    if s.startswith("Format setiap tabel mengikuti contoh ANSI"):
        i += 1
        continue
    m = re.fullmatch(r"\*\*Tabel 3\.\S* (.+)\*\*", s)
    if m:
        pending_caption = m.group(1)
        i += 1
        continue
    m = re.fullmatch(r"\*\*(UC-\d+): (.+)\*\*", s)
    if m:
        pending_caption = f"Use Case Description {m.group(1)} {m.group(2)}"
        i += 1
        continue
    if s.startswith("|"):
        rows, i = parse_table(i)
        header, data = rows[0], rows[1:]
        if header[:2] == ["Elemen", "Isi"]:
            pending_elem = {r[0]: r[1] for r in data}
            continue
        if header == ["Aktor", "Sistem"] and pending_elem is not None:
            caption("Tabel", pending_caption or pending_elem.get("Nama Use Case", ""))
            uc_table(pending_elem, data)
            pending_elem, pending_caption = None, None
            continue
        if pending_caption:
            caption("Tabel", pending_caption)
            pending_caption = None
        md_table(header, data)
        continue
    m = re.match(r"(\d+)\. (.+)", s)
    if m:
        list_par(m.group(1) + ".", m.group(2))
        i += 1
        continue
    if s.startswith("- "):
        list_par("•", s[2:])
        i += 1
        continue
    text = s.replace(" *(kalimat & sitasi bawaan template)*", "")
    if "Gambar 3.x berikut ini" in text:
        text = text.replace("Gambar 3.x", f"Gambar 3.{SEQ['Gambar'] + 1}")
        body_par(text)
        picture(XP_FIG if os.path.exists(XP_FIG) else xp_image, "Siklus Hidup Metode Pengembangan eXtreme Programming (Ambler, 2002)")
        i += 1
        continue
    if s.startswith("*") and s.endswith("*") and not s.startswith("**"):
        body_par(s.strip("*"), italic=True, indent=False)
    else:
        body_par(text)
    i += 1

# bibliography field with pre-rendered entries (Word regenerates them from the sources on update)
bp = new_par("Body Text")
bp.paragraph_format.left_indent, bp.paragraph_format.first_line_indent = Cm(1.0), Cm(-1.0)
bp.paragraph_format.line_spacing = 1.5
field(bp, "BIBLIOGRAPHY \\l 1057", "")
# replace the placeholder result run with formatted entries in following paragraphs is not possible inside one
# field result, so render the entries as the field result text separated by line breaks
res_t = bp._p.findall(".//" + qn("w:t"))[-1]
res_r = res_t.getparent()
res_r.remove(res_t)
for k, entry in enumerate(biblio):
    parts = re.split(r"(\*[^*]+\*)", entry)
    for part in parts:
        if not part:
            continue
        r = OxmlElement("w:r")
        if part.startswith("*") and part.endswith("*"):
            rPr = OxmlElement("w:rPr"); rPr.append(OxmlElement("w:i")); r.append(rPr); part = part[1:-1]
        t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = part; r.append(t)
        res_r.addprevious(r)
    if k < len(biblio) - 1:
        r = OxmlElement("w:r"); r.append(OxmlElement("w:br")); r.append(OxmlElement("w:br")); res_r.addprevious(r)
body.remove(anchor)

# ------------------------------------------------------------------ Word bibliography sources (customXml)
B = "http://schemas.openxmlformats.org/officeDocument/2006/bibliography"
TYPE = {"article": "JournalArticle", "book": "Book", "conference": "ConferenceProceedings", "report": "Report"}


def source_xml(r):
    x = [f"<b:Source><b:Tag>{r['tag']}</b:Tag><b:SourceType>{TYPE[r['type']]}</b:SourceType>",
         f"<b:Guid>{{{abs(hash(r['tag'])) % 10**8:08d}-0000-4000-8000-{abs(hash(r['title'])) % 10**12:012d}}}</b:Guid>"]
    if r.get("authors"):
        persons = "".join(f"<b:Person><b:Last>{escape(a[0])}</b:Last>" + (f"<b:First>{escape(a[1])}</b:First>" if a[1] else "")
                          + "</b:Person>" for a in r["authors"])
        x.append(f"<b:Author><b:Author><b:NameList>{persons}</b:NameList></b:Author></b:Author>")
    else:
        x.append(f"<b:Author><b:Author><b:Corporate>{escape(r['corporate'])}</b:Corporate></b:Author></b:Author>")
    x.append(f"<b:Title>{escape(r['title'])}</b:Title><b:Year>{r['year']}</b:Year>")
    if r["type"] == "article":
        x.append(f"<b:JournalName>{escape(r['journal'])}</b:JournalName>")
    if r["type"] == "conference":
        x.append(f"<b:ConferenceName>{escape(r['journal'])}</b:ConferenceName>")
    for k, tag in (("volume", "Volume"), ("issue", "Issue"), ("pages", "Pages"), ("publisher", "Publisher"),
                   ("city", "City"), ("doi", "DOI"), ("url", "URL"), ("isbn", "StandardNumber"), ("edition", "Edition")):
        if r.get(k):
            x.append(f"<b:{tag}>{escape(str(r[k]))}</b:{tag}>")
    x.append("</b:Source>")
    return "".join(x)


sources = (f'<?xml version="1.0" standalone="no"?><b:Sources SelectedStyle="\\APASixthEditionOfficeOnline.xsl" '
           f'StyleName="APA" Version="6" xmlns:b="{B}" xmlns="{B}">' + "".join(source_xml(r) for r in REFS) + "</b:Sources>")
done = False
for rel in doc.part.rels.values():
    if rel.reltype.endswith("/customXml"):
        part = rel.target_part
        if b"bibliography" in part.blob:
            part._blob = sources.encode("utf-8")
            done = True
if not done:
    raise SystemExit("no bibliography customXml part in base document")

# breathing room after headings (Heading 1: 18 pt, Heading 2: 10 pt)
for _p in doc.paragraphs:
    _n = _p.style.name if _p.style is not None else ""
    if _n == "Heading 1":
        _p.paragraph_format.space_after = Pt(18)
    elif _n == "Heading 2":
        _p.paragraph_format.space_before = Pt(14)
        _p.paragraph_format.space_after = Pt(10)
    elif _n == "Heading 3":
        _p.paragraph_format.space_before = Pt(10)
        _p.paragraph_format.space_after = Pt(6)

# Indonesian thesis convention: front-matter lists and the bibliography are single-spaced inside an entry,
# with a gap between entries; the bibliography has a 1 cm hanging indent and is left-aligned.
from docx.enum.style import WD_STYLE_TYPE
for _name in ("toc 1", "toc 2", "toc 3", "table of figures"):
    try:
        _pf = doc.styles[_name].paragraph_format
    except KeyError:
        continue
    _pf.line_spacing = 1.0
    _pf.space_before = Pt(0)
    _pf.space_after = Pt(6)
try:
    _bib = doc.styles["Bibliography"]
except KeyError:
    _bib = doc.styles.add_style("Bibliography", WD_STYLE_TYPE.PARAGRAPH)
    _bib.base_style = doc.styles["Normal"]
_bpf = _bib.paragraph_format
_bpf.line_spacing = 1.0
_bpf.space_before = Pt(0)
_bpf.space_after = Pt(10)
_bpf.left_indent = Cm(1.0)
_bpf.first_line_indent = Cm(-1.0)
_bpf.alignment = WD_ALIGN_PARAGRAPH.LEFT


# no em dashes anywhere in the finished document (the cover title comes from the base template)
for _t in doc.element.body.iter(qn("w:t")):
    if _t.text and "\u2014" in _t.text:
        _t.text = _t.text.replace(" \u2014 ", ": ").replace("\u2014", ", ")
_em = sum(1 for t in doc.element.body.iter(qn("w:t")) if t.text and "\u2014" in t.text)
print("em dash remaining:", _em)
doc.save(OUT)
print("saved", OUT, "tables", SEQ["Tabel"], "figures", SEQ["Gambar"])
